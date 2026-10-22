"""Read-only import audit. Private paths and normalized source text stay local."""
from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl

from .common import CODES, digest, now, save


def model(path):
    book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    items, issues = {}, []
    try:
        for row in book["五维模型定义"].iter_rows():
            code = row[1].value
            if not isinstance(code, str) or not re.fullmatch(r"[MEDIC][1-4]", code):
                continue
            if code in items:
                issues.append(f"duplicate:{code}")
            # C:name, D:definition, E-H:concept A-D, I:spacer, J-M:execution A-D.
            selected = list(row[2:8]) + list(row[9:13])
            values = [cell.value for cell in selected]
            if any(not isinstance(x, str) or not x.strip() for x in values):
                issues.append(f"missing_definition_or_anchor:{code}")
            if any(cell.data_type == "f" for cell in selected):
                issues.append(f"formula_requires_review:{code}")
            items[code] = dict(zip(
                ["name", "definition"] + [f"concept_{x}" for x in "ABCD"]
                + [f"execution_{x}" for x in "ABCD"], values
            ))
    finally:
        book.close()
    return items, issues


def matrix(path):
    book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    rows, issues = {}, []
    try:
        sheet = next(s for s in book if s.title.startswith("明细-"))
        for row in sheet.iter_rows(min_row=2):
            if not any(c.value is not None for c in row):
                continue
            job, code, grade, score = [row[i].value for i in (0, 6, 8, 9)]
            key = (job, code)
            if key in rows:
                issues.append(f"duplicate_row:{row[0].row}")
            if code not in CODES or grade not in tuple("ABCD") or score != "ABCD".index(grade) + 1:
                issues.append(f"invalid_recommendation:{row[0].row}")
            rows[key] = {"grade": grade, "score": score}
    finally:
        book.close()
    jobs = {key[0] for key in rows}
    for job in jobs:
        if {c for j, c in rows if j == job} != set(CODES):
            issues.append("incomplete_job_coverage")
    return rows, issues


def matrix_profiles(path):
    """Read repeated job labels without guessing classifications from keys."""
    book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    profiles, issues = {}, []
    try:
        sheet = next(s for s in book if s.title.startswith('明细-'))
        if [c.value for c in next(sheet.iter_rows(max_row=1))[:5]] != ['配置键','一级分类','二级分类','岗位名称','职级']:
            raise ValueError('matrix_profile_header_changed')
        for row in sheet.iter_rows(min_row=2):
            if not any(c.value is not None for c in row):
                continue
            cells = row[:5]
            values = [c.value for c in cells]
            if any(c.data_type=='f' for c in cells) or any(not isinstance(v,str) or not v.strip() for v in values):
                issues.append('invalid_job_profile_row:'+str(row[0].row))
                continue
            job, *labels = values
            profile = dict(zip(('family','track','title','level'),labels))
            if job in profiles and profiles[job] != profile:
                issues.append('conflicting_job_profile_row:'+str(row[0].row))
            else:
                profiles[job] = profile
    finally:
        book.close()
    return profiles, issues


def demo(path):
    text = Path(path).read_text(encoding="utf-8-sig")
    prefix = "window.MEDIC_DATA ="
    if not text.startswith(prefix):
        raise ValueError("unexpected_demo_format")
    return json.loads(text[len(prefix):].strip().removesuffix(";"))


def field_diff(before, after):
    return {code: [field for field in after[code] if before[code].get(field) != after[code][field]]
            for code in sorted(before.keys() & after.keys()) if before[code] != after[code]}


def dimensions(path):
    book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        return {row[0]: list(row[1:4]) for row in book["五维模型定义"].values
                if row[0] in tuple("MEDIC")}
    finally:
        book.close()


def audit(root, output):
    root = Path(root)
    paths = {
        "legacy_model": root / "核心胜任力模型260906_MEDIC.xlsx",
        "delivery_model": root / "交付0908/核心胜任力模型260906_MEDIC.xlsx",
        "matrix_260914": root / "岗位能力推荐值矩阵/MEDIC各岗位职级能力推荐值矩阵_260914_岗位职级更新版.xlsx",
        "demo_260908": root / "交付0908/MEDIC胜任力评估离线Demo/model-data.js",
        "policy_260922": root / "MEDIC胜任力评估_AI评估维度适用性判断_260922.md",
    }
    hashes = {name: digest(path.read_bytes()) for name, path in paths.items()}
    legacy, legacy_issues = model(paths["legacy_model"])
    latest, model_issues = model(paths["delivery_model"])
    old_dimensions, new_dimensions = dimensions(paths["legacy_model"]), dimensions(paths["delivery_model"])
    if set(new_dimensions) != set("MEDIC") or any(not all(v) for v in new_dimensions.values()):
        model_issues.append("invalid_dimension_definitions")
    recommendations, matrix_issues = matrix(paths["matrix_260914"])
    old = demo(paths["demo_260908"])
    old_model = {item["code"]: {"name": item["name"], "definition": item["definition"],
                  **{f"concept_{k}": v for k, v in item["concept"].items()},
                  **{f"execution_{k}": v for k, v in item["observations"].items()}}
                 for item in old["abilities"]}
    old_rows = {(r["key"], r["code"]): {"grade": r["grade"], "score": r["score"]}
                for r in old["recommendations"]}
    policy = paths["policy_260922"].read_text(encoding="utf-8-sig")
    groups = {}
    for number, mode in enumerate(("suggest", "clue_only", "human_only", "risk_scan"), 1):
        match = re.search(rf"\*\*第\s*{number}\s*组[^\n]*\n([^\n]+)", policy)
        groups[mode] = re.findall(r"[MEDIC][1-4]", match[1]) if match else []
    grouped = sum(groups.values(), [])
    policy_issues = [] if len(grouped) == 17 and set(grouped) == set(CODES) else ["invalid_policy_partition"]
    e3_line = next((line for line in policy.splitlines() if line.startswith("| E3 ")), "")
    conflicts = ([{"code": "E3", "table": "suggest", "summary": "clue_only",
                   "candidate_resolution": "clue_only", "requires_business_signoff": True}]
                 if "AI 初评" in e3_line and "E3" in groups["clue_only"] else [])
    jobs = {j for j, _ in recommendations}
    checks = {"model_17": set(latest) == set(CODES) and not model_issues,
              "matrix_49_833": len(jobs) == 49 and len(recommendations) == 833 and not matrix_issues,
              "policy_partition": not policy_issues}
    report = {
        "created_at": now(), "status": "pass" if all(checks.values()) else "fail",
        "publication_status": "awaiting_business_signoff", "source_hashes": hashes, "checks": checks,
        "counts": {"legacy_model": len(legacy), "delivery_model": len(latest),
                   "matrix_jobs": len(jobs), "matrix_rows": len(recommendations),
                   "demo_jobs": len({j for j, _ in old_rows}), "demo_rows": len(old_rows)},
        "legacy_missing_codes": sorted(latest.keys() - legacy.keys()),
        "dimension_definition_changes": [k for k in sorted(new_dimensions) if old_dimensions.get(k) != new_dimensions[k]],
        "legacy_changed_fields": field_diff(legacy, latest),
        "demo_changed_fields": field_diff(old_model, latest),
        "matrix_delta": {"added_pairs": len(recommendations.keys() - old_rows.keys()),
                         "removed_pairs": len(old_rows.keys() - recommendations.keys()),
                         "changed_common_pairs": sum(recommendations[k] != old_rows[k]
                                                     for k in recommendations.keys() & old_rows.keys())},
        "policy_groups": groups, "policy_conflicts": conflicts,
        "issues": {"legacy": legacy_issues, "model": model_issues, "matrix": matrix_issues, "policy": policy_issues},
    }
    save(output, report)
    return report
